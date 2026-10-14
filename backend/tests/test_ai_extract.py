"""AI 地基的单元测试：文档取文本 + 模型客户端配置校验。

不做真实网络调用——这里只验证「缺 Key 时报错」「解析结果对不对」，
真实调用留到接口联调时再测。
"""
import io

import pytest

from app.ai import AIError, _error_message, chat
from app.extract import ExtractError, extract_text


# ---------- extract：文件字节 -> 纯文本 ----------


def test_extract_txt_utf8():
    data = "2026级专业培养计划".encode("utf-8")
    assert extract_text("plan.txt", data) == "2026级专业培养计划"


def test_extract_txt_gbk():
    data = "教学大纲".encode("gb18030")
    assert extract_text("plan.txt", data) == "教学大纲"


def test_extract_docx_paragraph_and_table():
    import docx

    document = docx.Document()
    document.add_paragraph("专业培养计划")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "课程"
    table.rows[0].cells[1].text = "学分"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("plan.docx", buffer.getvalue())
    assert "专业培养计划" in text
    assert "课程" in text and "学分" in text


def test_extract_xlsx_sheet_and_rows():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "课程表"
    sheet.append(["课程", "学分"])
    sheet.append(["高等数学", 5])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = extract_text("plan.xlsx", buffer.getvalue())
    assert "课程表" in text
    assert "高等数学" in text


def test_extract_rejects_unsupported_format():
    with pytest.raises(ExtractError):
        extract_text("old.doc", b"whatever")


def test_extract_rejects_blank_content():
    with pytest.raises(ExtractError):
        extract_text("empty.txt", b"   ")


# ---------- ai：模型客户端 ----------


def test_chat_requires_api_key(app, monkeypatch):
    monkeypatch.delenv("COSMOS_API_KEY", raising=False)
    with app.app_context():
        app.config["COSMOS_API_KEY"] = ""
        with pytest.raises(AIError):
            chat([{"role": "user", "content": "你好"}])


def test_error_message_maps_platform_status():
    class FakeResponse:
        status_code = 401

        def json(self):
            return {"error": {"message": "invalid API key"}}

    message = _error_message(FakeResponse())
    assert "401" in message
    assert "invalid API key" in message
