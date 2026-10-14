"""把上传的文档转成纯文本，供大模型问答 / 预审使用。

项目里只存文件、不解析正文，这一步是 AI 功能的前置。解析库全部按需导入：
缺哪个库就返回一条带安装命令的提示，不影响应用启动。

目前支持：txt / md / csv（按文本读）、pdf（pdfplumber）、docx（python-docx）、
xlsx（openpyxl）。旧版二进制格式 doc / xls / ppt 不支持，需要先另存为新格式。
"""
from io import BytesIO
from pathlib import Path


class ExtractError(Exception):
    """文档解析失败。message 是可以直接展示给用户的中文说明。"""


TEXT_EXTENSIONS = ("txt", "md", "csv")
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS + ("pdf", "docx", "xlsx")


def extract_text(filename, data):
    """按扩展名把文件字节转成纯文本。"""
    if not data:
        raise ExtractError("文件内容为空")

    extension = Path(filename or "").suffix.lower().lstrip(".")
    if not extension:
        raise ExtractError("文件缺少扩展名，无法判断格式")

    handler = _HANDLERS.get(extension)
    if handler is None:
        supported = " / ".join(f".{item}" for item in SUPPORTED_EXTENSIONS)
        raise ExtractError(f"暂不支持解析 .{extension} 的正文，请换成 {supported}")

    try:
        text = handler(BytesIO(data))
    except ExtractError:
        raise
    except Exception as exc:  # 解析库抛的异常类型很杂，统一转成业务异常
        raise ExtractError(f"解析 .{extension} 文件失败：{exc}") from exc

    text = (text or "").strip()
    if not text:
        raise ExtractError(f".{extension} 文件里没有解析到文字内容")
    return text


def _decode_bytes(raw):
    for encoding in ("utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_text(stream):
    return _decode_bytes(stream.read())


def _read_pdf(stream):
    try:
        import pdfplumber
    except ImportError as exc:
        raise ExtractError(
            "缺少 pdfplumber，请执行 python -m pip install pdfplumber"
        ) from exc

    with pdfplumber.open(stream) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _read_docx(stream):
    try:
        import docx
    except ImportError as exc:
        raise ExtractError(
            "缺少 python-docx，请执行 python -m pip install python-docx"
        ) from exc

    document = docx.Document(stream)
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(blocks)


def _read_xlsx(stream):
    try:
        import openpyxl
    except ImportError as exc:
        raise ExtractError(
            "缺少 openpyxl，请执行 python -m pip install openpyxl"
        ) from exc

    workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    try:
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"[工作表] {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    lines.append("\t".join(cells))
        return "\n".join(lines)
    finally:
        workbook.close()


_HANDLERS = {
    **{extension: _read_text for extension in TEXT_EXTENSIONS},
    "pdf": _read_pdf,
    "docx": _read_docx,
    "xlsx": _read_xlsx,
}
