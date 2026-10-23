"""Excel/CSV table import and spreadsheet formula recalculation."""

from __future__ import annotations

import csv
import io
import re
from copy import deepcopy
from datetime import date, datetime, time
from pathlib import Path

from .formula_engine import (
    FormulaError,
    FormulaEvaluationError,
    FormulaEvaluator,
    FormulaSyntaxError,
    UnsupportedFormulaError,
    format_formula_result,
)


IMPORT_EXTENSIONS = {"xlsx", "xls", "csv"}
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_IMPORT_COLUMNS = 30
MAX_IMPORT_CELL_LENGTH = 5000
_CELL_RE = re.compile(r"^([A-Z]{1,3})(\d+)$")


class TableImportError(ValueError):
    def __init__(self, message: str, code: str = "invalid_import"):
        self.code = code
        super().__init__(message)


def cell_meta_key(row_key: str, column_key: str) -> str:
    return f"{row_key}:{column_key}"


def _value_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.15g}"
    return str(value)


def _column_letters(index: int) -> str:
    number = index + 1
    letters = ""
    while number > 0:
        remainder = (number - 1) % 26
        letters = chr(65 + remainder) + letters
        number = (number - 1) // 26
    return letters


def _split_cell_reference(reference: str) -> tuple[int, int]:
    match = _CELL_RE.match(reference.replace("$", "").upper())
    if not match:
        raise FormulaEvaluationError("#REF!")
    letters, row_text = match.groups()
    column = 0
    for char in letters:
        column = column * 26 + (ord(char) - 64)
    return int(row_text), column


def _column_number(reference: str) -> int:
    return _split_cell_reference(f"{reference}1")[1]


def remap_formula_references(formula, deleted_rows=None, deleted_columns=None):
    """Adjust A1 references after whole rows/columns are removed."""
    deleted_rows = sorted(set(deleted_rows or []))
    deleted_columns = sorted(set(deleted_columns or []))
    if not deleted_rows and not deleted_columns:
        return formula

    def remap(match):
        column_text = match.group(1)
        row_number = int(match.group(2))
        column_number = _column_number(column_text)
        if row_number in deleted_rows or column_number in deleted_columns:
            return "#REF!"
        row_shift = sum(1 for item in deleted_rows if item < row_number)
        column_shift = sum(1 for item in deleted_columns if item < column_number)
        return f"{_column_letters(column_number - column_shift - 1)}{row_number - row_shift}"

    parts = re.split(r'("(?:[^"]|"")*")', formula)
    for index in range(0, len(parts), 2):
        parts[index] = re.sub(
            r"\$?([A-Za-z]{1,3})\$?(\d+)",
            remap,
            parts[index],
        )
        parts[index] = re.sub(
            r"#REF!:[A-Za-z]{1,3}\d+|[A-Za-z]{1,3}\d+:#REF!",
            "#REF!",
            parts[index],
        )
        parts[index] = parts[index].replace("#REF!:#REF!", "#REF!")
    return "".join(parts)


class TableCalculator:
    """Recalculate formula cells against the current table grid."""

    def __init__(self, columns, rows, cell_meta, import_mode=False):
        self.columns = deepcopy(columns or [])
        self.rows = deepcopy(rows or [])
        self.cell_meta = deepcopy(cell_meta or {})
        self.import_mode = import_mode
        self.report = {
            "formula_count": 0,
            "recalculated_formula_count": 0,
            "locked_formulas": [],
            "formula_errors": [],
        }
        self._rows_by_key = {
            row.get("row_key"): row for row in self.rows if row.get("row_key")
        }
        self._column_index = {
            column.get("key"): index
            for index, column in enumerate(self.columns)
            if column.get("key")
        }
        self._coordinate = {}
        self._address = {}
        for row_number, row in enumerate(self.rows, start=2):
            for column_index, column in enumerate(self.columns):
                address = f"{_column_letters(column_index)}{row_number}"
                key = cell_meta_key(row.get("row_key"), column.get("key"))
                self._coordinate[address] = key
                self._address[key] = address
        self._evaluating = set()
        self._evaluated = {}
        self.evaluator = FormulaEvaluator(self)

    def cell(self, address: str):
        address = address.replace("$", "").upper()
        row_number, column_number = _split_cell_reference(address)
        if row_number == 1:
            if 1 <= column_number <= len(self.columns):
                return self.columns[column_number - 1].get("name", "")
            return 0.0

        key = self._coordinate.get(address)
        if key is None:
            return 0.0
        if key in self._evaluated:
            return self._evaluated[key]
        if key in self._evaluating:
            raise FormulaEvaluationError("#CIRC!")

        row_key, column_key = key.split(":", 1)
        row = self._rows_by_key.get(row_key)
        if row is None:
            return 0.0
        meta = self.cell_meta.get(key) or {}
        formula = meta.get("formula")
        if not formula:
            return (row.get("cells") or {}).get(column_key, "")
        if meta.get("locked"):
            return (row.get("cells") or {}).get(column_key, "")

        self._evaluating.add(key)
        try:
            value = self.evaluator.evaluate(formula)
        except UnsupportedFormulaError as exc:
            if not self.import_mode:
                self.report["formula_errors"].append(
                    {
                        "cell": address,
                        "formula": formula,
                        "error": str(exc),
                    }
                )
                raise
            cached = _value_text((row.get("cells") or {}).get(column_key, ""))
            meta["locked"] = True
            meta["warning"] = f"{exc}，已保留导入时的结果"
            self.cell_meta[key] = meta
            self._evaluated[key] = cached
            self.report["locked_formulas"].append(
                {
                    "cell": address,
                    "formula": formula,
                    "reason": str(exc),
                }
            )
            return cached
        except FormulaSyntaxError as exc:
            if not self.import_mode:
                self.report["formula_errors"].append(
                    {
                        "cell": address,
                        "formula": formula,
                        "error": f"公式语法错误：{exc}",
                    }
                )
                raise
            cached = _value_text((row.get("cells") or {}).get(column_key, ""))
            meta["locked"] = True
            meta["warning"] = f"公式无法解析：{exc}，已保留导入时的结果"
            self.cell_meta[key] = meta
            self._evaluated[key] = cached
            self.report["locked_formulas"].append(
                {
                    "cell": address,
                    "formula": formula,
                    "reason": f"公式无法解析：{exc}",
                }
            )
            return cached
        except FormulaEvaluationError as exc:
            value = exc.display
            self.report["formula_errors"].append(
                {
                    "cell": address,
                    "formula": formula,
                    "error": exc.display,
                }
            )
        finally:
            self._evaluating.discard(key)

        rendered = format_formula_result(value)
        if len(rendered) > MAX_IMPORT_CELL_LENGTH:
            rendered = "#VALUE!"
        cells = dict(row.get("cells") or {})
        cells[column_key] = rendered
        row["cells"] = cells
        self._evaluated[key] = value
        self.report["recalculated_formula_count"] += 1
        return value

    def range(self, start: str, end: str):
        start_row, start_column = _split_cell_reference(start)
        end_row, end_column = _split_cell_reference(end)
        if start_row > end_row:
            start_row, end_row = end_row, start_row
        if start_column > end_column:
            start_column, end_column = end_column, start_column
        values = []
        for row_number in range(start_row, end_row + 1):
            for column_number in range(start_column, end_column + 1):
                address = f"{_column_letters(column_number - 1)}{row_number}"
                values.append(self.cell(address))
        return values

    def recalculate(self):
        formula_keys = [
            key for key, meta in self.cell_meta.items() if meta.get("formula")
        ]
        self.report["formula_count"] = len(formula_keys)
        for key in formula_keys:
            address = self._address.get(key)
            if address:
                self.cell(address)
        return self.columns, self.rows, self.cell_meta, self.report


def calculate_grid(columns, rows, cell_meta, import_mode=False):
    calculator = TableCalculator(columns, rows, cell_meta, import_mode=import_mode)
    return calculator.recalculate()


def _next_unique_name(name: str, used: set[str]) -> str:
    candidate = name
    suffix = 2
    while candidate.casefold() in used:
        candidate = f"{name} ({suffix})"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _matrix_to_grid(matrix, formula_cells=None, cached_cells=None):
    formula_cells = formula_cells or {}
    cached_cells = cached_cells or {}
    nonempty = [
        (row, column)
        for row, values in enumerate(matrix)
        for column, value in enumerate(values)
        if value not in (None, "") or (row, column) in formula_cells
    ]
    if not nonempty:
        raise TableImportError("文件中没有可导入的数据", "empty_import")

    first_row = min(row for row, _ in nonempty)
    last_row = max(max(row for row, _ in nonempty), max((r for r, _ in formula_cells), default=0))
    first_column = min(column for _, column in nonempty)
    last_column = max(max(column for _, column in nonempty), max((c for _, c in formula_cells), default=0))
    column_count = last_column - first_column + 1
    data_row_count = last_row - first_row
    if column_count > MAX_IMPORT_COLUMNS:
        raise TableImportError(
            f"表格超过 {MAX_IMPORT_COLUMNS} 列，请精简后再导入",
            "too_many_columns",
        )
    if data_row_count > MAX_IMPORT_ROWS:
        raise TableImportError(
            f"表格超过 {MAX_IMPORT_ROWS} 行数据，请精简后再导入",
            "too_many_rows",
        )

    raw_headers = []
    for column in range(first_column, last_column + 1):
        value = formula_cells.get((first_row, column))
        if isinstance(value, str) and value.startswith("="):
            value = cached_cells.get((first_row, column))
        else:
            value = matrix[first_row][column] if column < len(matrix[first_row]) else ""
        raw_headers.append(_value_text(value).strip())

    used_names = set()
    columns = []
    for offset, raw_name in enumerate(raw_headers):
        fallback = f"列{_column_letters(offset)}"
        name = _next_unique_name(raw_name or fallback, used_names)
        columns.append({"key": f"col{offset + 1}", "name": name[:80], "type": "text"})

    rows = []
    cell_meta = {}
    for offset, row_number in enumerate(range(first_row + 1, last_row + 1)):
        row_key = f"r{offset + 1}"
        cells = {}
        for column_offset, column in enumerate(range(first_column, last_column + 1)):
            column_key = columns[column_offset]["key"]
            value = matrix[row_number][column] if (
                row_number < len(matrix) and column < len(matrix[row_number])
            ) else ""
            formula = formula_cells.get((row_number, column))
            if not formula and isinstance(value, str) and value.startswith("="):
                formula = value
            if formula:
                value = cached_cells.get((row_number, column), "")
            text = _value_text(value)
            if len(text) > MAX_IMPORT_CELL_LENGTH:
                raise TableImportError(
                    f"{_column_letters(column)}{row_number + 1} 单元格超过 {MAX_IMPORT_CELL_LENGTH} 个字符",
                    "cell_too_long",
                )
            cells[column_key] = text
            if formula:
                cell_meta[cell_meta_key(row_key, column_key)] = {
                    "formula": formula,
                    "locked": False,
                }
        rows.append({"row_key": row_key, "cells": cells})

    columns, rows, cell_meta, report = calculate_grid(
        columns, rows, cell_meta, import_mode=True
    )
    missing_results = []
    for item in report["locked_formulas"]:
        address = item["cell"]
        row_number, column_number = _split_cell_reference(address)
        if row_number < 2:
            continue
        row = rows[row_number - 2] if row_number - 2 < len(rows) else None
        column = columns[column_number - 1] if column_number - 1 < len(columns) else None
        if row and column and not (row.get("cells") or {}).get(column.get("key"), ""):
            missing_results.append(item)
    if missing_results:
        examples = "、".join(
            f"{item['cell']}（{item['reason']}）" for item in missing_results[:5]
        )
        raise TableImportError(
            f"以下公式没有可保留的计算结果，已拒绝导入：{examples}",
            "missing_formula_result",
        )
    return columns, rows, cell_meta, report


def _read_xlsx(data: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise TableImportError("服务器缺少 Excel 解析组件", "xlsx_unavailable") from exc

    try:
        formula_book = load_workbook(
            io.BytesIO(data), read_only=False, data_only=False
        )
        value_book = load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    except Exception as exc:
        raise TableImportError(
            "无法读取 Excel 文件，请确认文件未损坏或加密",
            "invalid_excel",
        ) from exc

    if getattr(formula_book, "_external_links", None):
        raise TableImportError("暂不支持带外部链接的 Excel 文件", "external_links")

    worksheet = formula_book.worksheets[0]
    value_worksheet = value_book[worksheet.title]
    actual_cells = [
        cell
        for cell in worksheet._cells.values()
        if cell.value not in (None, "")
    ]
    if not actual_cells:
        raise TableImportError("文件中没有可导入的数据", "empty_import")

    first_row = min(cell.row for cell in actual_cells)
    last_row = max(cell.row for cell in actual_cells)
    first_column = min(cell.column for cell in actual_cells)
    last_column = max(cell.column for cell in actual_cells)
    if last_row - first_row > MAX_IMPORT_ROWS:
        raise TableImportError(
            f"表格超过 {MAX_IMPORT_ROWS} 行数据，请精简后再导入",
            "too_many_rows",
        )
    if last_column - first_column + 1 > max(MAX_IMPORT_COLUMNS * 10, 300):
        raise TableImportError("工作表有效列过多，请精简后再导入", "too_many_columns")

    matrix = []
    formula_cells = {}
    cached_cells = {}
    for row_number in range(first_row, last_row + 1):
        row_values = []
        for column_number in range(first_column, last_column + 1):
            cell = worksheet.cell(row=row_number, column=column_number)
            value = cell.value
            row_values.append(value)
            if isinstance(value, str) and value.startswith("="):
                formula_cells[(row_number - first_row, column_number - first_column)] = value
                cached_cells[(row_number - first_row, column_number - first_column)] = (
                    value_worksheet.cell(row=row_number, column=column_number).value
                )
        matrix.append(row_values)
    return _matrix_to_grid(matrix, formula_cells, cached_cells), worksheet.title


def _read_xls(data: bytes):
    try:
        import xlrd
    except ImportError as exc:
        raise TableImportError("服务器缺少旧版 Excel 解析组件", "xls_unavailable") from exc

    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        worksheet = workbook.sheet_by_index(0)
    except Exception as exc:
        raise TableImportError(
            "无法读取旧版 Excel 文件，请确认文件未损坏或加密",
            "invalid_excel",
        ) from exc

    matrix = [
        [worksheet.cell_value(row, column) for column in range(worksheet.ncols)]
        for row in range(worksheet.nrows)
    ]
    result = _matrix_to_grid(matrix)
    return result, worksheet.name


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise TableImportError("CSV 编码无法识别，请另存为 UTF-8 后重试", "invalid_csv")


def _read_csv(data: bytes):
    text = _decode_csv(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    matrix = list(csv.reader(io.StringIO(text), dialect))
    result = _matrix_to_grid(matrix)
    return result, "CSV"


def parse_import_file(filename: str, data: bytes):
    extension = Path(filename or "").suffix.lower().lstrip(".")
    if extension not in IMPORT_EXTENSIONS:
        raise TableImportError(
            "仅支持 .xlsx、.xls、.csv 文件",
            "invalid_file_type",
        )
    if not data:
        raise TableImportError("请选择要导入的文件", "empty_file")
    if len(data) > MAX_IMPORT_BYTES:
        raise TableImportError("导入文件不能超过 5 MB", "file_too_large")

    if extension == "xlsx":
        (columns, rows, cell_meta, report), sheet_name = _read_xlsx(data)
    elif extension == "xls":
        (columns, rows, cell_meta, report), sheet_name = _read_xls(data)
    else:
        (columns, rows, cell_meta, report), sheet_name = _read_csv(data)

    return {
        "columns": columns,
        "rows": rows,
        "cell_meta": cell_meta,
        "sheet_name": sheet_name,
        "extension": extension,
        "report": report,
    }
