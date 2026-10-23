r"""Generate the repeatable teaching-document review scenario.

Run from the repository root:

    .\.venv\Scripts\python.exe tools\generate_ai_review_scenario.py

The generated DOCX files are committed because they are also seed data for the
application and upload material for the AI review smoke test.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "backend" / "demo_data" / "ai_review_scenario"

INK = RGBColor(0, 0, 0)
MUTED = RGBColor(80, 80, 80)
HEADER_FILL = "1F4E78"
PALE_FILL = "EAF1F8"
LIGHT_BORDER = "D9D9D9"


def set_run_font(run, name="宋体", size=10.5, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.45):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LIGHT_BORDER)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cell_text(cell, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    set_paragraph_spacing(paragraph, before=0, after=0, line=1.25)
    run = paragraph.add_run(str(text))
    set_run_font(run, size=9.3, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths_cm, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    header_row = table.rows[0]
    repeat_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = Cm(widths_cm[index])
        shade_cell(cell, HEADER_FILL)
        set_cell_text(
            cell,
            header,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=font_size, bold=True, color=RGBColor(255, 255, 255))

    for row_index, values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Cm(widths_cm[col_index])
            if row_index % 2 == 1:
                shade_cell(cell, PALE_FILL)
            align = (
                WD_ALIGN_PARAGRAPH.CENTER
                if col_index == 0 or len(str(value)) <= 8
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            set_cell_text(cell, value, align=align)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=font_size)
    return table


def add_body(doc, text, *, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    set_paragraph_spacing(paragraph, after=7, line=1.55)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet_list(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.first_line_indent = Cm(-0.37)
        set_paragraph_spacing(paragraph, after=3, line=1.4)
        run = paragraph.add_run(item)
        set_run_font(run)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    set_paragraph_spacing(
        paragraph,
        before=14 if level == 1 else 9,
        after=6,
        line=1.15,
    )
    run = paragraph.add_run(text)
    set_run_font(
        run,
        name="黑体",
        size=14 if level == 1 else 11.5,
        bold=True,
        color=INK,
    )
    return paragraph


def add_title(doc, title, subtitle):
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=12, after=5, line=1.15)
    for index, part in enumerate(title.split("\n")):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(part)
        set_run_font(run, name="黑体", size=20, bold=True, color=INK)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, after=13, line=1.2)
    run = paragraph.add_run(subtitle)
    set_run_font(run, size=9.5, color=MUTED)


def add_metadata_table(doc, metadata):
    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for index, (label, value) in enumerate(metadata):
        label_cell, value_cell = table.rows[index].cells
        label_cell.width = Cm(3.2)
        value_cell.width = Cm(12.8)
        shade_cell(label_cell, PALE_FILL)
        set_cell_text(label_cell, label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(value_cell, value)
    doc.add_paragraph()


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "黑体"
        style.font.color.rgb = INK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        if style._element.pPr is not None:
            border = style._element.pPr.find(qn("w:pBdr"))
            if border is not None:
                style._element.pPr.remove(border)


def new_doc():
    doc = Document()
    configure_document(doc)
    return doc


def add_page_break(doc):
    doc.add_section(WD_SECTION.NEW_PAGE)


def build_program_plan():
    doc = new_doc()
    add_title(
        doc,
        "2026级数据科学与大数据技术专业培养方案",
        "专业代码 080910T｜适用年级 2026 级｜文件版本 V1.0",
    )
    add_metadata_table(
        doc,
        [
            ("培养单位", "计算机与人工智能学院"),
            ("修业年限", "四年，授予工学学士学位"),
            ("总学分", "160 学分"),
            ("方案状态", "现行有效"),
        ],
    )
    add_body(
        doc,
        "本方案规定数据科学与大数据技术专业 2026 级学生四年的培养目标、毕业要求、课程结构和学分要求。方案用于指导教师组织教学、学生规划学习进度，也作为课程教学大纲审核的上位依据。",
    )

    add_heading(doc, "一、培养目标")
    add_body(
        doc,
        "毕业生能够面向数据治理、数据分析与智能应用等领域，综合运用数学、统计和计算机知识解决工程问题；能够在多学科团队中承担数据工程、模型开发或项目管理工作，并遵守工程伦理和数据安全规范。",
    )
    add_bullet_list(
        doc,
        [
            "毕业五年左右能够独立承担数据工程项目中的核心开发或分析任务。",
            "能够针对业务问题选择合适的数据处理、存储和分析方法。",
            "能够在团队中清晰表达技术方案，并对实施结果承担责任。",
        ],
    )

    add_heading(doc, "二、毕业要求与指标点")
    add_body(
        doc,
        "毕业要求采用学校 2026 版工程教育认证通用框架。课程教学大纲必须标明课程支撑的指标点，并提供教学内容与考核证据。",
    )
    add_table(
        doc,
        ["毕业要求", "指标点编号", "指标点摘要"],
        [
            ["工程知识", "1.2", "能够运用计算机与数学知识表达数据工程问题"],
            ["问题分析", "2.3", "能够建立概念模型、关系模型和约束方案"],
            ["设计/开发解决方案", "3.2", "能够设计满足业务约束的数据管理与处理方案"],
            ["研究", "4.1", "能够基于实验或实测数据形成有效结论"],
            ["使用现代工具", "5.2", "能够选择并使用数据库、开发框架和分析工具"],
            ["工程与社会", "6.1", "能够评价专业实践对安全、隐私和社会的影响"],
            ["个人和团队", "9.2", "能够在团队中承担设计、开发、测试或汇报职责"],
        ],
        [3.2, 2.7, 10.1],
    )

    add_heading(doc, "三、学分与学时结构")
    add_table(
        doc,
        ["课程模块", "最低学分", "学时占比", "主要构成"],
        [
            ["通识教育", "50", "31.3%", "思想政治、外语、体育、通识选修"],
            ["学科基础", "40", "25.0%", "数学、统计、程序设计"],
            ["专业核心", "35", "21.9%", "数据库、数据挖掘、分布式计算"],
            ["专业选修", "25", "15.6%", "机器学习、数据治理、可视化"],
            ["集中实践", "10", "6.2%", "课程设计、专业实习、毕业设计"],
            ["合计", "160", "100%", "毕业最低学分要求"],
        ],
        [4.0, 3.1, 3.2, 5.7],
    )

    add_heading(doc, "四、数据库原理及应用课程定位")
    add_body(
        doc,
        "数据库原理及应用为专业核心课，安排在第三学期，4 学分、64 学时。课程承担关系数据模型、SQL、数据库设计、事务与恢复、查询优化等核心内容，并为后续的数据仓库、数据治理与毕业设计提供基础。",
    )
    add_table(
        doc,
        ["课程信息", "内容"],
        [
            ["课程代码", "DSB201"],
            ["先修课程", "程序设计基础、离散数学"],
            ["后修课程", "数据仓库与 BI、大数据平台实践"],
            ["考核结构", "过程性考核 50%，终结性考核 50%"],
            ["审核要求", "课程目标、教学内容、考核方式与支撑指标点须形成完整闭环"],
        ],
        [4.0, 12.0],
    )
    return doc


def build_indicator_map():
    doc = new_doc()
    add_title(
        doc,
        "数据科学与大数据技术专业\n毕业要求指标点 2026版",
        "文件版本 V1.0｜适用对象 2026 级｜教学大纲审核参考文件",
    )
    add_body(
        doc,
        "本文件将数据科学与大数据技术专业的毕业要求拆分为可观察、可考核的指标点。课程教学大纲中的课程目标应至少支撑一个指标点，并在教学内容与考核证据中形成对应关系。",
    )
    add_heading(doc, "一、指标点清单")
    add_table(
        doc,
        ["编号", "指标点表述", "课程可提供的证据"],
        [
            ["1.2", "能够运用数学、统计与计算机基础知识，对数据工程问题进行规范表达。", "概念模型、推导过程、建模作业"],
            ["2.3", "能够识别数据管理需求，建立概念模型、关系模型与完整性约束。", "ER 图、关系模式、约束说明"],
            ["3.2", "能够设计满足性能、安全与业务约束的数据管理和处理方案。", "数据库设计报告、SQL 实现、方案评审"],
            ["4.1", "能够设计实验，采集和分析数据，并形成有依据的结论。", "实验记录、查询结果分析、优化对比"],
            ["5.2", "能够选择并使用数据库系统、开发框架和数据分析工具完成复杂任务。", "工具操作记录、项目成果、版本材料"],
            ["6.1", "能够评价数据管理实践对隐私、安全和社会责任的影响。", "安全设计、伦理讨论、风险分析"],
            ["9.2", "能够在多学科团队中承担设计、开发、测试或项目汇报职责。", "分工记录、项目汇报、同伴评价"],
        ],
        [1.6, 9.8, 4.6],
    )
    doc.add_page_break()

    add_heading(doc, "二、课程映射要求")
    add_bullet_list(
        doc,
        [
            "每门专业核心课原则上支撑 2 至 4 个指标点。",
            "每个课程目标只标注一个主支撑指标点，避免一张表重复覆盖全部要求。",
            "课程目标中的能力动词应为可评价的行为，例如解释、设计、分析、实现、评价。",
            "团队协作指标点必须通过小组项目、分工记录或同伴评价取得证据，不能只依赖闭卷考试。",
            "指标点编号须与培养方案保持一致，临时编号或未列入方案的编号不能作为有效映射。",
        ],
    )

    add_heading(doc, "三、数据库原理及应用课程建议映射")
    add_table(
        doc,
        ["课程目标方向", "建议指标点", "最低证据要求"],
        [
            ["理解关系模型与 SQL", "2.3", "建模作业、SQL 实验"],
            ["完成数据库应用设计", "3.2", "课程项目设计文档与实现"],
            ["分析事务、索引和优化", "5.2", "实验报告与查询执行计划"],
            ["参与团队项目", "9.2", "分工记录、项目汇报、同伴评价"],
        ],
        [5.8, 3.0, 7.2],
    )
    return doc


def build_writing_spec():
    doc = new_doc()
    add_title(
        doc,
        "数据库原理及应用课程\n教学大纲编写规范 2026版",
        "文件版本 V1.0｜发布单位 计算机与人工智能学院｜审核日期 2026年9月",
    )
    add_body(
        doc,
        "本规范规定数据库原理及应用课程教学大纲的结构、目标映射、学时计算、考核设计和版本说明。大纲只有同时满足必填内容、数据一致和考核可追溯三项要求，才能提交学院审核。",
    )

    add_heading(doc, "一、适用范围与基本要求")
    add_bullet_list(
        doc,
        [
            "适用于数据科学与大数据技术专业开设的数据库原理及应用课程。",
            "文件标题应包含课程名称和文件类型，正文应写明课程代码、学分、总学时、适用对象和先修课程。",
            "课程目标须与毕业要求指标点逐项对应，每项目标至少有一条教学内容依据和一条考核依据。",
            "数据、术语、指标点编号和课程名称在全文必须一致。",
        ],
    )

    add_heading(doc, "二、必填章节")
    add_table(
        doc,
        ["章节", "必须包含的信息", "审核要点"],
        [
            ["基本信息", "课程代码、学分、总学时、学期、适用专业、先修课程", "字段齐全且与培养方案一致"],
            ["课程目标", "目标编号、能力表述、支撑指标点", "目标可评价、指标点有效"],
            ["教学内容与学时", "教学单元、主要内容、学时、目标对应关系", "单元学时合计等于总学时"],
            ["教学方式", "课堂讲授、实验、项目及各自组织方式", "与能力目标匹配"],
            ["考核方式", "考核项目、比例、评价标准、对应目标", "比例合计 100%"],
            ["教材与资源", "教材、参考书、数据库平台或实验环境", "能够支撑课程实施"],
            ["版本说明", "版本号、编制人、审核记录、主要变更", "可追溯修订内容"],
        ],
        [3.4, 8.4, 5.2],
    )

    add_heading(doc, "三、课程目标与指标点映射")
    add_body(
        doc,
        "课程目标应使用可观察的行为描述。团队协作、沟通表达等能力目标必须安排项目汇报、同伴评价或过程记录等可留存证据的考核方式。仅设置闭卷考试，不足以证明团队协作目标已经达成。",
    )
    add_body(
        doc,
        "课程目标编号、指标点编号、教学内容编号和考核项目编号应保持一致。例如目标 D2 对应指标点 3.2，则教学内容和考核依据也应明确标注 D2，便于审核人员追踪。",
    )

    add_heading(doc, "四、学时计算")
    add_body(
        doc,
        "总学时是课程实施的总量，应等于理论教学、实验教学、项目实践等各单元学时之和。审核时须展示参与计算的全部数字，不得只写“共 64 学时”而不拆分明细。出现差额时，应先核对单元学时和总学时，再修改课程安排。",
    )
    add_table(
        doc,
        ["核对项", "计算规则", "合格示例"],
        [
            ["总学时", "各教学单元学时之和", "8 + 20 + 12 + 10 + 8 + 6 = 64"],
            ["理论学时", "各讲授单元学时之和", "48"],
            ["实验学时", "各实验或项目单元学时之和", "16"],
            ["一致性", "总学时 = 理论学时 + 实验学时", "64 = 48 + 16"],
        ],
        [3.4, 7.8, 5.8],
    )

    add_heading(doc, "五、考核方式")
    add_body(
        doc,
        "过程性考核和终结性考核比例之和必须为 100%。过程性考核原则上不低于 40%，终结性考核不高于 60%。过程性考核应覆盖实验、项目、课堂参与或阶段性成果，终结性考核应验证课程核心知识掌握情况。",
    )
    add_table(
        doc,
        ["考核项目", "建议比例", "应形成的证据", "对应目标"],
        [
            ["课堂参与", "10%", "考勤与课堂任务记录", "D1"],
            ["基础实验", "20%", "SQL 实验报告与结果", "D1、D2"],
            ["综合项目", "20%", "需求、设计、实现、汇报、分工记录", "D3、D4"],
            ["期末闭卷考试", "50%", "试卷与评分标准", "D1-D3"],
            ["合计", "100%", "所有比例相加必须为 100%", "-"],
        ],
        [3.6, 2.7, 7.5, 3.2],
    )

    add_heading(doc, "六、审核闭环与版本说明")
    add_body(
        doc,
        "提交审核的新版本须在变更说明中写明上一版问题及处理结果。人工退回后，新版应能够逐项说明“已解决、未解决或无法判断”；修订内容必须能够与总学时、考核比例和目标映射等原始依据对应。",
    )

    add_heading(doc, "七、合格性检查表")
    add_table(
        doc,
        ["序号", "检查项", "通过条件"],
        [
            ["1", "核心内容完整", "基本信息、目标、内容、学时、考核和资源均有实质内容"],
            ["2", "关键数据一致", "课程名称、代码、学分、学时和指标点编号全文一致"],
            ["3", "版本变化清晰", "首版说明初始状态；修订版说明主要变更并关联退回意见"],
            ["4", "结论有材料依据", "问题和修订结论引用正文数字或规范要求"],
            ["5", "必填章节齐全", "七类必填信息没有缺失"],
            ["6", "总学时一致", "总学时等于各教学单元学时之和"],
            ["7", "目标有内容支撑", "每项目标至少对应一个教学单元"],
            ["8", "目标有考核支撑", "每项目标至少对应一个考核项目，团队目标含协作证据"],
            ["9", "考核比例合计 100%", "过程性考核与终结性考核加总为 100%"],
            ["10", "模板字段符合要求", "标题、课程代码、学分、学时等固定字段可识别"],
        ],
        [1.2, 4.7, 10.7],
    )
    return doc


def add_draft_body(doc, corrected=False):
    if corrected:
        add_metadata_table(
            doc,
            [
                ("课程代码", "DSB201"),
                ("课程学分", "4 学分"),
                ("总学时", "64 学时"),
                ("适用学期", "2026-2027 学年第一学期"),
                ("适用专业", "数据科学与大数据技术"),
                ("先修课程", "程序设计基础、离散数学"),
                ("编制人", "胡军成"),
                ("审核记录", "2026-10-23 根据课程组退回意见完成整改"),
                ("版本状态", "整改稿 V2.0，提交复审"),
            ],
        )
    else:
        add_metadata_table(
            doc,
            [
                ("课程代码", "DSB201"),
                ("课程学分", "4 学分"),
                ("总学时", "72 学时"),
                ("适用学期", "2026-2027 学年第一学期"),
                ("适用专业", "数据科学与大数据技术"),
                ("先修课程", "程序设计基础、离散数学"),
                ("编制人", "胡军成"),
                ("审核记录", "2026-09-28 提交课程组首次审核"),
                ("版本状态", "首次提交，待审核"),
            ],
        )

    add_heading(doc, "一、课程简介")
    add_body(
        doc,
        "本课程面向数据科学与大数据技术专业本科生，系统讲解关系数据模型、结构化查询语言、数据库设计与规范化、事务并发控制和查询优化。课程通过实验和综合项目训练学生设计、实现和评价数据管理方案的能力。",
    )

    add_heading(doc, "二、课程目标与指标点")
    if corrected:
        objectives = [
            ["D1", "能够解释关系模型、完整性约束和 SQL 的核心概念。", "2.3", "关系模型与 SQL"],
            ["D2", "能够依据业务需求完成概念设计、关系模式设计和规范化。", "3.2", "数据库设计与规范化"],
            ["D3", "能够分析事务并发、恢复、索引与查询优化问题。", "5.2", "事务、索引与优化实验"],
            ["D4", "能够在团队项目中承担数据库设计、实现或汇报职责，并依据分工记录完成同伴评价。", "9.2", "综合项目"],
        ]
    else:
        objectives = [
            ["D1", "能够解释关系模型、完整性约束和 SQL 的核心概念。", "2.3", "关系模型与 SQL"],
            ["D2", "能够依据业务需求完成概念设计、关系模式设计和规范化。", "3.2", "数据库设计与规范化"],
            ["D3", "能够分析事务并发、恢复、索引与查询优化问题。", "5.2", "事务、索引与优化实验"],
            ["D4", "能够以团队形式完成数据库应用项目并清晰表达设计结果。", "9.2", "综合项目"],
        ]
    add_table(
        doc,
        ["目标编号", "课程目标", "支撑指标点", "主要教学内容"],
        objectives,
        [1.8, 8.7, 2.6, 4.9],
    )

    add_heading(doc, "三、教学内容与学时")
    add_table(
        doc,
        ["教学单元", "主要内容", "理论", "实验", "合计", "目标"],
        [
            ["数据库系统概述", "数据库体系结构、数据模型、数据管理系统", "8", "0", "8", "D1"],
            ["关系模型与 SQL", "关系代数、完整性约束、DDL、DML、查询", "16", "4", "20", "D1、D2"],
            ["数据库设计与规范化", "ER 模型、关系模式转换、函数依赖、范式", "8", "4", "12", "D2"],
            ["事务并发与恢复", "事务、锁、隔离级别、日志与恢复", "10", "0", "10", "D3"],
            ["索引与查询优化", "B+ 树、执行计划、索引选择、性能分析", "4", "4", "8", "D3"],
            ["综合项目", "需求分析、数据库实现、测试、汇报与同伴评价", "2", "4", "6", "D2、D4"],
            ["合计", "各教学单元学时之和", "48", "16", "64", "-"],
        ],
        [4.0, 7.2, 1.7, 1.7, 1.7, 2.0],
    )
    if corrected:
        add_body(
            doc,
            "学时核对：8 + 20 + 12 + 10 + 8 + 6 = 64 学时；理论 48 学时，实验 16 学时，总学时与理论、实验学时之和一致。",
        )

    add_heading(doc, "四、教学方式")
    add_body(
        doc,
        "课程以课堂讲授建立概念框架，以数据库实验训练 SQL 和性能分析能力，以综合项目训练数据库设计和协作能力。实验使用统一数据库环境，项目按 3 至 4 人分组实施，保留需求文档、设计稿、代码、演示记录和分工证据。",
    )

    add_heading(doc, "五、考核方式")
    if corrected:
        assessment = [
            ["课堂参与", "10%", "考勤与课堂任务", "D1"],
            ["基础实验", "20%", "SQL 与优化实验报告", "D1、D2、D3"],
            ["综合项目", "20%", "设计文档、实现、项目汇报、同伴评价", "D2、D4"],
            ["期末闭卷考试", "50%", "概念理解、建模与 SQL 分析", "D1-D3"],
            ["合计", "100%", "过程性考核 50%，终结性考核 50%", "-"],
        ]
    else:
        assessment = [
            ["课堂参与", "10%", "考勤记录", "D1"],
            ["基础实验", "20%", "SQL 与优化实验报告", "D1、D2、D3"],
            ["期末闭卷考试", "60%", "概念理解、建模与 SQL 分析", "D1-D3"],
            ["合计", "90%", "过程性考核 30%，终结性考核 60%", "-"],
        ]
    add_table(
        doc,
        ["考核项目", "比例", "评价依据", "对应目标"],
        assessment,
        [3.8, 2.2, 8.4, 3.6],
    )

    if corrected:
        add_body(
            doc,
            "目标 D4 的考核证据包括小组项目报告、现场汇报、组内分工记录和同伴评价；同伴评价得分单独记录，并由教师核定。考核比例核对：10% + 20% + 20% + 50% = 100%。",
        )

    add_heading(doc, "六、教材与教学资源")
    add_bullet_list(
        doc,
        [
            "主要教材：数据库系统概论，高等教育出版社。",
            "参考资源：PostgreSQL 15 官方文档、课程实验指导书、公开数据集。",
            "实验环境：PostgreSQL 15、DBeaver、Python 3.11 与 psycopg。",
        ],
    )

    if corrected:
        add_heading(doc, "七、修订说明")
        add_table(
            doc,
            ["上版问题", "本版处理", "核验依据"],
            [
                ["总学时 72 与单元合计 64 不一致", "将总学时统一为 64", "教学内容表合计与学时核对式"],
                ["过程性考核只有 30%，比例合计 90%", "增加 20% 综合项目，比例调整为 50% + 50%", "考核方式表合计 100%"],
                ["目标 D4 缺少团队协作考核", "增加项目汇报、分工记录和同伴评价", "综合项目评价依据"],
            ],
            [6.1, 5.8, 5.1],
        )
    return doc


def build_draft(corrected=False):
    doc = new_doc()
    title = "数据库原理及应用课程教学大纲"
    subtitle = (
        "文件版本 V2.0｜整改复审稿｜计算机与人工智能学院"
        if corrected
        else "文件版本 V1.0｜首次提交审核｜计算机与人工智能学院"
    )
    add_title(doc, title, subtitle)
    add_draft_body(doc, corrected=corrected)
    return doc


def build_assessment_report():
    doc = new_doc()
    add_title(
        doc,
        "2025-2026学年第二学期\n数据库原理及应用课程考核分析报告",
        "课程代码 DSB201｜授课对象 2024 级｜报告状态 已归档",
    )
    add_metadata_table(
        doc,
        [
            ("学生总数", "120 人"),
            ("课程学分", "4 学分"),
            ("总评及格率", "95.0%"),
            ("总评平均分", "78.6 分"),
        ],
    )
    add_body(
        doc,
        "本学期课程按课堂参与 10%、基础实验 20%、综合项目 20% 和期末闭卷考试 50% 评定成绩。课程目标 D4 通过项目汇报和同伴评价取得证据。本报告用于分析目标达成情况并改进下一轮教学。",
    )

    add_heading(doc, "一、成绩分布")
    add_table(
        doc,
        ["分数段", "人数", "占比", "累计占比"],
        [
            ["90-100", "18", "15.0%", "15.0%"],
            ["80-89", "36", "30.0%", "45.0%"],
            ["70-79", "40", "33.3%", "78.3%"],
            ["60-69", "20", "16.7%", "95.0%"],
            ["60 以下", "6", "5.0%", "100%"],
            ["合计", "120", "100%", "-"],
        ],
        [3.0, 3.0, 3.0, 3.0],
    )

    add_heading(doc, "二、课程目标达成情况")
    add_table(
        doc,
        ["目标", "支撑指标点", "达成度", "主要依据"],
        [
            ["D1", "2.3", "0.84", "期末考试概念题、SQL 基础实验"],
            ["D2", "3.2", "0.79", "数据库设计项目、期末考试设计题"],
            ["D3", "5.2", "0.76", "索引与查询优化实验"],
            ["D4", "9.2", "0.81", "项目成果、现场汇报、同伴评价"],
        ],
        [2.0, 3.0, 3.0, 8.2],
    )

    add_heading(doc, "三、主要问题与改进")
    add_body(
        doc,
        "查询优化实验的达成度最低。学生在解释执行计划时，能够识别全表扫描，但对索引选择与代价估算的分析不充分。下一轮将增加执行计划对比实验，并要求学生在报告中说明索引变化对查询代价的影响。",
        bold_lead="查询优化实验的达成度最低。",
    )
    add_body(
        doc,
        "团队项目总体完成情况良好，但少数小组的分工记录晚于开发过程。后续将把分工记录纳入第一次阶段检查，确保过程证据及时留存。",
        bold_lead="团队项目总体完成情况良好",
    )
    add_heading(doc, "四、材料归档")
    add_bullet_list(
        doc,
        [
            "成绩单、试卷分析表、项目评分表和同伴评价表已归档。",
            "课程目标达成度计算表随报告一并保存。",
            "下一轮教学大纲修订应继续保留过程性考核 50%、终结性考核 50% 的结构。",
        ],
    )
    return doc


def save_doc(doc, filename):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / filename
    doc.save(path)
    return path


def main():
    outputs = [
        save_doc(build_program_plan(), "01_2026级数据科学与大数据技术专业培养方案.docx"),
        save_doc(build_indicator_map(), "02_数据科学与大数据技术专业毕业要求指标点_2026版.docx"),
        save_doc(build_writing_spec(), "03_数据库原理及应用课程教学大纲编写规范_2026版.docx"),
        save_doc(build_draft(corrected=False), "04_数据库原理及应用课程教学大纲_待审核_v1.docx"),
        save_doc(build_draft(corrected=True), "05_数据库原理及应用课程教学大纲_整改稿_v2.docx"),
        save_doc(build_assessment_report(), "06_数据库原理及应用课程考核分析报告.docx"),
    ]
    for path in outputs:
        print(f"generated: {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
